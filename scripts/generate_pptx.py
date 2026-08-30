import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank slide

    # Tactical Color Palette
    COLOR_BG = RGBColor(7, 11, 20)       # #070b14
    COLOR_SURFACE = RGBColor(15, 23, 42) # #0f172a
    COLOR_TEXT = RGBColor(241, 245, 249) # #f1f5f9
    COLOR_MUTED = RGBColor(148, 163, 184)# #94a3b8
    COLOR_CYAN = RGBColor(6, 182, 212)   # #06b6d4
    COLOR_BLUE = RGBColor(59, 130, 246)  # #3b82f6
    COLOR_EMERALD = RGBColor(16, 185, 129)# #10b981
    COLOR_AMBER = RGBColor(245, 158, 11) # #f59e0b
    COLOR_RED = RGBColor(239, 68, 68)    # #ef4444

    def add_tactical_slide(title, subtitle=None):
        slide = prs.slides.add_slide(blank_layout)
        
        # Background rect
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()

        # Top Accent Line
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.04))
        accent.fill.solid()
        accent.fill.fore_color.rgb = COLOR_CYAN
        accent.line.fill.background()

        # Title Text Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle.upper()
            p2.font.size = Pt(11)
            p2.font.bold = True
            p2.font.color.rgb = COLOR_CYAN
            p2.space_before = Pt(4)

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.text = "IBVAP  |  INTELLIGENT BORDER VIDEO ANALYTICS PLATFORM  |  SIH INTERNAL ROUND"
        fp.font.size = Pt(9)
        fp.font.color.rgb = COLOR_MUTED

        return slide

    def add_card(slide, left, top, width, height, title, items, accent_color=COLOR_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_SURFACE
        card.line.color.rgb = accent_color
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = accent_color
        p.space_after = Pt(8)

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"•  {item}"
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = COLOR_TEXT
            p_item.space_after = Pt(6)

    # ----------------------------------------------------
    # SLIDE 1: Title Slide
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG
    bg1.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "IBVAP"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p = tf1.add_paragraph()
    p.text = "INTELLIGENT BORDER VIDEO ANALYTICS PLATFORM"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT
    p.space_before = Pt(8)

    p = tf1.add_paragraph()
    p.text = "Autonomous AI Surveillance, Deterministic Geofencing & Cryptographic Chain of Custody"
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_MUTED
    p.space_before = Pt(12)

    p = tf1.add_paragraph()
    p.text = "Smart India Hackathon (SIH) Internal Round  |  Verified & Demo Ready  |  198+ FPS Throughput"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_EMERALD
    p.space_before = Pt(24)

    # ----------------------------------------------------
    # SLIDE 2: Problem Statement
    # ----------------------------------------------------
    s2 = add_tactical_slide("The Border Surveillance Challenge", "Operational Vulnerabilities in Traditional CCTV")
    add_card(s2, 0.8, 1.8, 5.6, 4.8, "Conventional CCTV Limitations", [
        "Passive Recording: Records crimes post-facto rather than intercepting in real time.",
        "Cognitive Operator Fatigue: Detection accuracy drops >70% after 20 minutes.",
        "Overwhelming False Alarms: Simple motion sensors trigger on wind, rain, and shadows.",
        "Scalability Bottleneck: Impossible to monitor hundreds of feeds simultaneously."
    ], COLOR_RED)
    add_card(s2, 6.9, 1.8, 5.6, 4.8, "Evidence & Forensic Deficits", [
        "Lack of Chain of Custody: Raw video files on disk lack cryptographic proof.",
        "Vulnerable to Tampering: Internal actors can alter or delete frames undetected.",
        "Inadmissible Evidence: Repudiation risk in military and legal tribunals.",
        "Slow Forensic Retrieval: Hours spent searching unindexed video archives."
    ], COLOR_AMBER)

    # ----------------------------------------------------
    # SLIDE 3: Proposed Solution
    # ----------------------------------------------------
    s3 = add_tactical_slide("The IBVAP Autonomous Solution", "Closed-Loop AI Intelligence & Cryptographic Custody")
    add_card(s3, 0.8, 1.8, 3.6, 4.8, "1. Autonomous AI Vision", [
        "YOLOv8 Real-Time Inference",
        "ByteTrack Identity Association",
        "Ground Foot-Point Projection",
        "Arbitrary Polygon Geofencing"
    ], COLOR_CYAN)
    add_card(s3, 4.8, 1.8, 3.6, 4.8, "2. Real-Time Alerting", [
        "Deterministic State Machine",
        "Zero Alert Spam / Flooding",
        "Sub-10ms WebSocket Push",
        "Tactical Dark Command Center"
    ], COLOR_BLUE)
    add_card(s3, 8.8, 1.8, 3.6, 4.8, "3. Verifiable Forensics", [
        "Server-Side SHA-256 Hashing",
        "On-Demand Tamper Detection",
        "Bit-Level Mismatch Flagging",
        "Immutable PostgreSQL Audit Log"
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 4: System Architecture
    # ----------------------------------------------------
    s4 = add_tactical_slide("End-to-End System Architecture", "Decoupled High-Throughput Surveillance Pipeline")
    add_card(s4, 0.8, 1.8, 5.6, 4.8, "Edge AI & Vision Layer", [
        "Camera Feed: RTSP / Webcam / Video Stream",
        "YOLOv8: Real-Time Bounding Box & Class Inference",
        "ByteTrack: Kalman Filter Trajectory Tracking",
        "Polygon Engine: Shapely Point-in-Polygon Ray Casting",
        "EventDispatcher: Scoped JWT Authenticated HTTP POST"
    ], COLOR_CYAN)
    add_card(s4, 6.9, 1.8, 5.6, 4.8, "Platform & Command Center", [
        "FastAPI Backend: REST API (/api/v1) + RBAC Auth",
        "PostgreSQL DB: Events, Alerts, Evidence & Audit Logs",
        "WebSocket Manager: Authenticated Push Handshake",
        "Next.js Tactical UI: Live Monitoring & Audio Synthesizer",
        "Integrity Service: Server-Authoritative SHA-256 Verification"
    ], COLOR_BLUE)

    # ----------------------------------------------------
    # SLIDE 5: AI Computer Vision Pipeline
    # ----------------------------------------------------
    s5 = add_tactical_slide("AI Computer Vision Pipeline", "YOLOv8 + ByteTrack + Polygon Ray-Casting")
    add_card(s5, 0.8, 1.8, 3.6, 4.8, "Detection (YOLOv8)", [
        "High-throughput object classification.",
        "Person, vehicle, and animal labels.",
        "Confidence threshold filtering (0.50+).",
        "Fail-closed model SHA-256 integrity."
    ], COLOR_CYAN)
    add_card(s5, 4.8, 1.8, 3.6, 4.8, "Tracking (ByteTrack)", [
        "Maintains persistent Track IDs.",
        "Associates low and high score boxes.",
        "Survives occlusion & temporary crossing.",
        "Kalman filter trajectory smoothing."
    ], COLOR_BLUE)
    add_card(s5, 8.8, 1.8, 3.6, 4.8, "Geofencing (Polygon PIP)", [
        "Normalized 2D coordinates [0.0 - 1.0].",
        "Bottom-center foot point extraction.",
        "Shapely ray-casting containment check.",
        "Resolution & aspect-ratio independent."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 6: Intrusion State Machine
    # ----------------------------------------------------
    s6 = add_tactical_slide("Intrusion Decision State Machine", "Eliminating False Positive Flooding")
    add_card(s6, 0.8, 1.8, 5.6, 4.8, "State Transition Logic", [
        "OUTSIDE ➔ INSIDE: State machine detects crossing and emits exactly ONE intrusion event.",
        "INSIDE (Continuous): Suppresses duplicate alarms while intruder remains in the zone.",
        "INSIDE ➔ OUTSIDE: Resets track state upon boundary exit.",
        "Hysteresis Buffer: Prevents jitter on polygon perimeter boundary lines."
    ], COLOR_AMBER)
    add_card(s6, 6.9, 1.8, 5.6, 4.8, "Operational Benefits", [
        "Zero Command Fatigue: 1 incident = 1 actionable alert card.",
        "Deterministic Auditability: Every state change logged with UTC timestamp.",
        "Idempotent Event IDs: Duplicate network packets return HTTP 409 Conflict.",
        "Bandwidth Efficient: Eliminates redundant database writes."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 7: Defense-in-Depth Security Matrix
    # ----------------------------------------------------
    s7 = add_tactical_slide("Defense-in-Depth Security Architecture", "Threat-Modeled Protection Across 15 Attack Vectors")
    add_card(s7, 0.8, 1.8, 3.6, 4.8, "Identity & API Access", [
        "Argon2id password hashing.",
        "15-Minute Scoped JWT tokens.",
        "TOTP RFC-6238 Multi-Factor Auth.",
        "Sliding-window brute force limiting."
    ], COLOR_CYAN)
    add_card(s7, 4.8, 1.8, 3.6, 4.8, "Authorization & RBAC", [
        "ADMINISTRATOR: Full configuration.",
        "OPERATOR: Alert ACK & monitoring.",
        "ANALYST: Read-only telemetry.",
        "AUDITOR: Immutable audit inspection."
    ], COLOR_BLUE)
    add_card(s7, 8.8, 1.8, 3.6, 4.8, "Integrity & Anti-Tamper", [
        "Model SHA-256 (Fail-Closed).",
        "Safe path traversal validation.",
        "Server-side SHA-256 evidence hashing.",
        "Append-only PostgreSQL audit trail."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 8: Cryptographic Evidence Vault
    # ----------------------------------------------------
    s8 = add_tactical_slide("Cryptographic Evidence & Chain of Custody", "Mathematical Proof of Zero Evidence Modification")
    add_card(s8, 0.8, 1.8, 5.6, 4.8, "Capture & Hashing Lifecycle", [
        "1. Instant Capture: Raw JPEG frame written to disk at exact moment of intrusion.",
        "2. Server-Side Hashing: 64-char SHA-256 digest computed immediately via hashlib.",
        "3. Atomic Persistence: Hash stored in PostgreSQL evidence table with UTC timestamp.",
        "4. Non-Repudiation: Immutable record links Event ➔ Alert ➔ Evidence."
    ], COLOR_CYAN)
    add_card(s8, 6.9, 1.8, 5.6, 4.8, "Authoritative Verification States", [
        "🟢 VERIFIED: Disk binary matches database hash digest (0 diff - untampered).",
        "🔴 TAMPER DETECTED: 1-bit file change flags exact SHA-256 mismatch.",
        "🟡 NOT_HASHED: Exposes one-click server-side hash generation.",
        "Audit Trail: Verification action logged in audit_logs with actor ID."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 9: Real-Time WebSocket Notification Feed
    # ----------------------------------------------------
    s9 = add_tactical_slide("Real-Time Streaming & WebSocket Architecture", "Sub-10ms Push Notifications Without Polling")
    add_card(s9, 0.8, 1.8, 5.6, 4.8, "Streaming Architecture", [
        "Initial REST Hydration: Loads historical alerts from PostgreSQL on mount.",
        "Authenticated WS Channel: ws://localhost:8000/api/v1/ws/events?token=<JWT>",
        "Post-Commit Broadcast: Notifications emitted only AFTER database commit.",
        "Single Global Socket: Managed in AlertProvider to prevent redundant pooling."
    ], COLOR_BLUE)
    add_card(s9, 6.9, 1.8, 5.6, 4.8, "Reliability & Resilience", [
        "Handshake Security: Invalid/expired JWT rejected with code 1008.",
        "Duplicate Prevention: Deduplicates cards on alert_id and event_id.",
        "Auto-Reconnect: Bounded exponential backoff (max 10s) on network drop.",
        "Native Audio Synthesizer: Immediate audible warning chime on incursion."
    ], COLOR_CYAN)

    # ----------------------------------------------------
    # SLIDE 10: Tactical Command Dashboard
    # ----------------------------------------------------
    s10 = add_tactical_slide("Operational Command Center UI", "Mission-Critical Tactical Interface (Next.js 14)")
    add_card(s10, 0.8, 1.8, 5.6, 4.8, "Command Center Features", [
        "Dominant Surveillance Viewport: Live feed with active polygon zone overlay.",
        "Tactical Threat Stack: Live alert queue with severity and confidence tags.",
        "System Telemetry Badge: Real-time BACKEND ● CONNECTED | WS | DB status.",
        "Interactive Evidence Modal: High-res snapshot viewer and SHA-256 verifier."
    ], COLOR_CYAN)
    add_card(s10, 6.9, 1.8, 5.6, 4.8, "Engineering Standards", [
        "Zero Production Mocks: All routes consume live /api/v1 backend endpoints.",
        "Dark Theme Aesthetics: High-contrast military slate & neon accents.",
        "Type-Safe TypeScript: Strict component contracts and shared schemas.",
        "Production Static Build: 13 optimized static pages compiled cleanly."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 11: Empirical Performance & Benchmarks
    # ----------------------------------------------------
    s11 = add_tactical_slide("Empirical Performance Benchmarks", "Verified Measurements on 1280x720 Benchmark Dataset")
    add_card(s11, 0.8, 1.8, 5.6, 4.8, "Measured Metrics", [
        "AI Video Ingestion: 198.58 FPS (300 frames @ 1280x720 video).",
        "Event Persistence: 50 – 65 ms detection-to-database commit.",
        "WebSocket Delivery: < 10 ms push latency to browser UI.",
        "Evidence SHA-256 Hashing: < 2 ms binary recalculation."
    ], COLOR_EMERALD)
    add_card(s11, 6.9, 1.8, 5.6, 4.8, "Regression Suite Results", [
        "Backend & AI Regression Suite: 183 / 183 PASSED (100% in 11.96s).",
        "Security Threat Audit Suite: 7 / 7 Threat Tests PASSED (0.61s).",
        "Frontend Code Quality: ESLint PASS (0 errors), TSC PASS (0 errors).",
        "Demo Rehearsal: Completed in 5 min 25 sec with 0 interventions."
    ], COLOR_CYAN)

    # ----------------------------------------------------
    # SLIDE 12: Failure Recovery & Resilience
    # ----------------------------------------------------
    s12 = add_tactical_slide("Failure Recovery & Robustness", "Graceful Degradation Without Synthetic Fallbacks")
    add_card(s12, 0.8, 1.8, 5.6, 4.8, "Failure Modes & Responses", [
        "Backend Offline: UI switches to BACKEND ● OFFLINE; no fake data.",
        "Backend Restart: Client auto-reconnects and resumes live monitoring.",
        "WebSocket Drop: Reconnects via exponential backoff (max 10s).",
        "Database Drop: Health check flags DB: DISCONNECTED."
    ], COLOR_AMBER)
    add_card(s12, 6.9, 1.8, 5.6, 4.8, "Security Failure Mitigations", [
        "Corrupted YOLO Model: Fail-closed RuntimeError halts engine.",
        "Tampered Evidence File: Flags 🔴 TAMPER DETECTED immediately.",
        "Duplicate Event Payload: HTTP 409 Conflict suppresses duplicate alarms.",
        "Path Traversal Attempt: Path resolver rejects arbitrary file access."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 13: Core Differentiation
    # ----------------------------------------------------
    s13 = add_tactical_slide("Core Differentiation: CCTV vs. IBVAP", "Transforming Dumb Cameras into Autonomous Security Nodes")
    add_card(s13, 0.8, 1.8, 5.6, 4.8, "Conventional CCTV", [
        "Passive video recording.",
        "Continuous human visual fatigue.",
        "Unbounded false alarm rates.",
        "Unauthenticated raw video files.",
        "No audit trail of operator actions."
    ], COLOR_RED)
    add_card(s13, 6.9, 1.8, 5.6, 4.8, "IBVAP Smart Platform", [
        "Autonomous YOLOv8 detection & ByteTrack tracking.",
        "Shapely polygon geofencing & state-machine filtering.",
        "Sub-10ms authenticated WebSocket alert delivery.",
        "Server-authoritative SHA-256 tamper verification.",
        "Append-only immutable PostgreSQL audit logging."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 14: Scalability & Enterprise Roadmap
    # ----------------------------------------------------
    s14 = add_tactical_slide("Scalability & Cloud Architecture Roadmap", "Transition from Localhost MVP to Enterprise Cluster")
    add_card(s14, 0.8, 1.8, 5.6, 4.8, "Phase 1: Current MVP Stack", [
        "Localhost Native Execution (Python + Node.js).",
        "Single-node PostgreSQL relational database.",
        "In-process async WebSocket manager.",
        "Local structured evidence filesystem."
    ], COLOR_BLUE)
    add_card(s14, 6.9, 1.8, 5.6, 4.8, "Phase 2: Enterprise Cloud Scale", [
        "Multi-Camera Kubernetes AI Worker Cluster (GPU TensorRT).",
        "Apache Kafka / Redis Streams distributed event broker.",
        "S3 / MinIO Object Storage with WORM tamper-proof policy.",
        "Distributed Redis Pub/Sub WebSocket gateway."
    ], COLOR_CYAN)

    # ----------------------------------------------------
    # SLIDE 15: Live Demonstration Sequence
    # ----------------------------------------------------
    s15 = add_tactical_slide("Live Demonstration Flow", "5-Minute Tactical Demonstration Walkthrough")
    add_card(s15, 0.8, 1.8, 11.733, 4.8, "Step-by-Step Tactical Sequence", [
        "1. Login & Dashboard: Authenticate as operator ➔ Show BACKEND ● CONNECTED | WS | DB.",
        "2. Surveillance & Polygon: Display active restricted polygon geofence overlay.",
        "3. Live Incursion: Run video/camera ➔ Person enters zone ➔ Instant alert banner & chime.",
        "4. Operator Acknowledgment: Click Acknowledge ➔ Status updates to ACKNOWLEDGED via PATCH.",
        "5. Evidence Vault: Open modal ➔ Show high-res snapshot & 64-char SHA-256 digest.",
        "6. Tamper Demonstration: Verify 🟢 (MATCH) ➔ Modify byte ➔ Verify 🔴 (MISMATCH) ➔ Restore 🟢.",
        "7. Audit Trail: Inspect /audit-logs ➔ Show EVIDENCE_VERIFIED and ALERT_ACKNOWLEDGED entries."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 16: Operational & Defense Impact
    # ----------------------------------------------------
    s16 = add_tactical_slide("Operational & Defense Impact", "High-Impact Technology for National Security")
    add_card(s16, 0.8, 1.8, 5.6, 4.8, "Tactical Advantages", [
        "Sub-Second Threat Awareness: Faster tactical interception of perimeter breaches.",
        "Eliminates Human Blind Spots: 24/7 continuous autonomous vigilance.",
        "Courtroom Admissibility: Cryptographic proof of custody for legal proceedings.",
        "Cost-Effective Retrofit: Works over existing commodity IP camera networks."
    ], COLOR_CYAN)
    add_card(s16, 6.9, 1.8, 5.6, 4.8, "SIH Value Proposition", [
        "Solves authentic Ministry of Home Affairs / Defense border problem statement.",
        "100% verified functional MVP with zero mock data.",
        "High performance (198+ FPS) and rigorous test coverage (183 tests).",
        "Clean, maintainable, and modular open architecture."
    ], COLOR_EMERALD)

    # ----------------------------------------------------
    # SLIDE 17: Development Timeline & Horizons
    # ----------------------------------------------------
    s17 = add_tactical_slide("Development Roadmap & Future Horizons", "From MVP Foundation to Advanced Sensor Fusion")
    add_card(s17, 0.8, 1.8, 5.6, 4.8, "Completed Milestones (M1 – M3.6)", [
        "M1: FastAPI REST platform + Argon2id Auth + PostgreSQL.",
        "M2: YOLOv8 + ByteTrack + Polygon PIP + Intrusion Engine.",
        "M3: Security hardening + WebSocket + Next.js Tactical UI + Forensics.",
        "Status: 183/183 Tests Passing | FROZEN & READY."
    ], COLOR_BLUE)
    add_card(s17, 6.9, 1.8, 5.6, 4.8, "Future Research Horizons", [
        "Multi-Spectral Thermal IR Camera Fusion for zero-light night operations.",
        "Cross-Camera Re-Identification (ReID) across multiple border towers.",
        "Autonomous PTZ Camera Slew-to-Cue tracking integration.",
        "Edge TensorRT deployment on solar-powered border poles."
    ], COLOR_CYAN)

    # ----------------------------------------------------
    # SLIDE 18: Conclusion & Q&A
    # ----------------------------------------------------
    s18 = prs.slides.add_slide(blank_layout)
    bg18 = s18.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg18.fill.solid()
    bg18.fill.fore_color.rgb = COLOR_BG
    bg18.line.fill.background()

    tb18 = s18.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.5))
    tf18 = tb18.text_frame
    tf18.word_wrap = True

    p = tf18.paragraphs[0]
    p.text = "IBVAP"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p = tf18.add_paragraph()
    p.text = "INTELLIGENT BORDER VIDEO ANALYTICS PLATFORM"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT
    p.space_before = Pt(6)

    p = tf18.add_paragraph()
    p.text = "From raw video pixels to verified, auditable operational intelligence."
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_MUTED
    p.space_before = Pt(12)

    p = tf18.add_paragraph()
    p.text = "🟢 183/183 Tests Passing  |  198+ FPS Throughput  |  FROZEN FOR SIH EVALUATION"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_EMERALD
    p.space_before = Pt(20)

    p = tf18.add_paragraph()
    p.text = "Thank you. We are ready for your questions."
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p.space_before = Pt(24)

    output_path = "SIH_FINAL_PRESENTATION.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
